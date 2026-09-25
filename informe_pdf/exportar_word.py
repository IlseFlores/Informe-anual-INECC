"""Exporta el Informe Anual de Calidad del Aire a un .docx editable.

Usa exactamente el mismo contenido que generar_informe.py (mismo texto, mismas
tablas y mismas gráficas), de modo que el Word y el PDF salgan siempre de la
misma fuente: los textos y las tablas quedan como texto/tablas de Word, y las
gráficas y mapas como imágenes.

Uso:  python exportar_word.py --anio 2024
"""
import argparse
import re
import tempfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.graphics import renderPDF
from reportlab.graphics.shapes import Drawing
from reportlab.platypus import Image as FlowImage
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, Spacer, Table

import generar_informe as g

# Fuente del Word: Calibri se ve igual en cualquier equipo. Si tienes Montserrat estática
# instalada, puedes cambiarla aquí (la variable de Windows se ve muy delgada en Word).
FUENTE = "Calibri"
NAVY = RGBColor(0x17, 0x3D, 0x4C)
TEXTO = RGBColor(0x32, 0x4D, 0x59)
MUTED = RGBColor(0x6C, 0x88, 0x94)
ANCHO_UTIL_PT = g.CONTENT_WIDTH


# ---------------------------------------------------------------------------
# Utilidades de formato
# ---------------------------------------------------------------------------
def _hex(color):
    return "%02X%02X%02X" % (round(color.red * 255), round(color.green * 255), round(color.blue * 255))


def _es_oscuro(color):
    lum = 0.299 * color.red + 0.587 * color.green + 0.114 * color.blue
    return lum < 0.62


def _fuente(run, tam=None, negrita=None, italica=None, color=None):
    run.font.name = FUENTE
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    if tam:
        run.font.size = Pt(tam)
    if negrita is not None:
        run.font.bold = negrita
    if italica is not None:
        run.font.italic = italica
    if color is not None:
        run.font.color.rgb = color


def _sombrear(celda, hex_color):
    tc_pr = celda._tc.get_or_add_tcPr()
    for viejo in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(viejo)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _bordes_tabla(tabla, color="D7DEE1"):
    tbl_pr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        bordes.append(b)
    tbl_pr.append(bordes)


def _campo_pagina(parrafo):
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)
    _fuente(run, 8, color=MUTED)


# ---------------------------------------------------------------------------
# Texto con marcado mínimo de reportlab (<b>, <i>, <br/>, <font ...>)
# ---------------------------------------------------------------------------
_TAG = re.compile(r"(<[^>]+>)")


def _agregar_runs(parrafo, marcado, tam, color, negrita_base=False, italica_base=False):
    negrita, italica = negrita_base, italica_base
    color_actual = [color]
    for trozo in _TAG.split(marcado):
        if not trozo:
            continue
        if trozo.startswith("<"):
            etiqueta = trozo.lower()
            if etiqueta in ("<b>", "<strong>"):
                negrita = True
            elif etiqueta in ("</b>", "</strong>"):
                negrita = negrita_base
            elif etiqueta == "<i>":
                italica = True
            elif etiqueta == "</i>":
                italica = italica_base
            elif etiqueta.startswith("<br"):
                parrafo.add_run().add_break()
            elif etiqueta.startswith("<font"):
                m = re.search(r"color=['\"]?#([0-9a-fA-F]{6})", trozo)
                if m:
                    color_actual.append(RGBColor.from_string(m.group(1).upper()))
            elif etiqueta == "</font>" and len(color_actual) > 1:
                color_actual.pop()
            continue
        texto = (trozo.replace("&nbsp;", " ").replace("&amp;", "&")
                 .replace("&lt;", "<").replace("&gt;", ">"))
        run = parrafo.add_run(texto)
        _fuente(run, tam, negrita, italica, color_actual[-1])


# ---------------------------------------------------------------------------
# Conversión de flowables
# ---------------------------------------------------------------------------
class Exportador:
    def __init__(self, anio):
        self.anio = anio
        self.doc = Document()
        self.tmp = Path(tempfile.mkdtemp(prefix="informe_word_"))
        self.n_img = 0
        self._preparar_documento()

    def _preparar_documento(self):
        seccion = self.doc.sections[0]
        seccion.orientation = WD_ORIENT.PORTRAIT
        seccion.page_width, seccion.page_height = Cm(21.59), Cm(27.94)
        seccion.left_margin = seccion.right_margin = Cm(2.2)
        seccion.top_margin, seccion.bottom_margin = Cm(2.4), Cm(2.0)

        normal = self.doc.styles["Normal"]
        normal.font.name = FUENTE
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        for nombre, tam in (("Heading 1", 15), ("Heading 2", 12)):
            estilo = self.doc.styles[nombre]
            estilo.font.name = FUENTE
            estilo.font.size = Pt(tam)
            estilo.font.bold = True
            estilo.font.color.rgb = NAVY
            estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
            estilo.element.rPr.rFonts.set(qn("w:ascii"), FUENTE)
            estilo.element.rPr.rFonts.set(qn("w:hAnsi"), FUENTE)

        encabezado = seccion.header.paragraphs[0]
        _agregar_runs(encabezado, f"SIMAJ · JALISCO  |  Informe Anual de Calidad del Aire {self.anio}", 8, MUTED)
        pie = seccion.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _agregar_runs(pie, "Sistema de Monitoreo Atmosférico de Jalisco (SIMAJ)  ·  Página ", 8, MUTED)
        _campo_pagina(pie)

    # -- portada --------------------------------------------------------
    def portada(self, datos):
        ruta = Path(datos.get("imagen_portada", ""))
        if ruta.exists():
            self.doc.add_picture(str(ruta), width=Cm(15.5))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for texto, tam, color in (
            ("I N F O R M E   A N U A L", 13, RGBColor(0x00, 0xC4, 0xB4)),
            ("Calidad del Aire", 34, NAVY),
            (f"— Jalisco {self.anio} —", 18, RGBColor(0x00, 0xC4, 0xB4)),
        ):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _fuente(p.add_run(texto), tam, True, None, color)
        p = self.doc.add_paragraph()
        _fuente(p.add_run(f"{datos.get('entidad_editora', '')}  ·  {datos.get('fecha_publicacion', '')}"), 9, None, None, MUTED)
        self.doc.add_page_break()

    # -- flowables -----------------------------------------------------
    def agregar(self, flowable):
        if isinstance(flowable, PageBreak):
            self.doc.add_page_break()
        elif isinstance(flowable, Spacer):
            return
        elif isinstance(flowable, KeepTogether):
            for f in flowable._content:
                self.agregar(f)
        elif isinstance(flowable, Paragraph):
            self._parrafo(flowable)
        elif isinstance(flowable, Table):
            self._tabla(flowable)
        elif isinstance(flowable, Drawing):
            self._dibujo(flowable)
        elif isinstance(flowable, g.MosaicoMapas):
            self._dibujo(flowable.dibujo_completo)
        elif isinstance(flowable, FlowImage):
            self.doc.add_picture(str(flowable.filename), width=Pt(min(flowable.drawWidth, ANCHO_UTIL_PT)))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif isinstance(flowable, list):
            for f in flowable:
                self.agregar(f)
        else:
            print(f"Aviso: elemento no exportado ({type(flowable).__name__})")

    def _parrafo(self, para):
        nombre = para.style.name
        texto = para.text
        if nombre == "h2":
            p = self.doc.add_paragraph(style="Heading 1")
            _agregar_runs(p, texto, 15, NAVY, negrita_base=True)
        elif nombre == "h3":
            p = self.doc.add_paragraph(style="Heading 2")
            _agregar_runs(p, texto, 12, NAVY, negrita_base=True)
        elif nombre in ("tabla_caption", "caption"):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            _agregar_runs(p, texto, 9.5, NAVY, negrita_base=True)
        elif nombre == "credito_nombre":
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            _agregar_runs(p, texto, 13, NAVY, negrita_base=True)
        elif nombre == "credito_cargo":
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(14)
            _agregar_runs(p, texto, 10.5, MUTED)
        elif nombre == "credito_equipo":
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            _agregar_runs(p, texto, 11.5, NAVY, negrita_base=True)
        elif nombre == "bibliografia":
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(16)
            p.paragraph_format.first_line_indent = Pt(-16)
            p.paragraph_format.space_after = Pt(6)
            _agregar_runs(p, "- " + texto, 10, TEXTO)
        elif nombre in ("nota", "nota_destacada"):
            p = self.doc.add_paragraph()
            _agregar_runs(p, texto, 8.5, MUTED)
        else:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(8)
            _agregar_runs(p, texto, 11, TEXTO)

    def _dibujo(self, dibujo):
        self.n_img += 1
        pdf = self.tmp / f"d{self.n_img}.pdf"
        png = self.tmp / f"d{self.n_img}.png"
        # Algunos dibujos (p. ej. leyendas) se extienden un poco más allá de su
        # ancho declarado: se amplía la caja para que no se recorten.
        ancho_original = dibujo.width
        dibujo.width = max(ancho_original, dibujo.getBounds()[2] + 2)
        try:
            renderPDF.drawToFile(dibujo, str(pdf))
        finally:
            ancho_real, dibujo.width = dibujo.width, ancho_original
        pagina = fitz.open(str(pdf))[0]
        pagina.get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False).save(str(png))
        ancho = min(ancho_real, ANCHO_UTIL_PT)
        self.doc.add_picture(str(png), width=Pt(ancho))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # -- tablas --------------------------------------------------------
    @staticmethod
    def _norm(coord, n):
        return coord if coord >= 0 else n + coord

    def _tabla(self, tabla):
        filas = tabla._cellvalues
        n_filas, n_cols = len(filas), len(filas[0])
        docx_tabla = self.doc.add_table(rows=n_filas, cols=n_cols)
        docx_tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        docx_tabla.autofit = False
        _bordes_tabla(docx_tabla)

        total = sum(tabla._colWidths)
        ancho_tabla = min(total, ANCHO_UTIL_PT)
        escala = ancho_tabla / total
        anchos = [w * escala for w in tabla._colWidths]

        # Fondos (BACKGROUND / ROWBACKGROUNDS)
        fondos = {}
        for cmd in tabla._bkgrndcmds:
            op, (sc, sr), (ec, er) = cmd[0], cmd[1], cmd[2]
            sc, ec = self._norm(sc, n_cols), self._norm(ec, n_cols)
            sr, er = self._norm(sr, n_filas), self._norm(er, n_filas)
            if op == "BACKGROUND":
                for r in range(sr, er + 1):
                    for c in range(sc, ec + 1):
                        fondos[(r, c)] = cmd[3]
            elif op == "ROWBACKGROUNDS":
                colores = cmd[3]
                for i, r in enumerate(range(sr, er + 1)):
                    for c in range(sc, ec + 1):
                        fondos[(r, c)] = colores[i % len(colores)]

        tam = 6.6 if n_cols > 9 else 8.4
        for r, fila in enumerate(filas):
            for c, contenido in enumerate(fila):
                celda = docx_tabla.cell(r, c)
                celda.width = Pt(anchos[c])
                color_fondo = fondos.get((r, c))
                texto_claro = False
                if color_fondo is not None and hasattr(color_fondo, "red"):
                    _sombrear(celda, _hex(color_fondo))
                    texto_claro = _es_oscuro(color_fondo)
                self._llenar_celda(celda, contenido, tam, texto_claro)

        # Encabezado repetido y sin separarse de la tabla
        rep = tabla.repeatRows
        n_enc = (max(rep) + 1 if isinstance(rep, (tuple, list)) and rep else int(rep or 0))
        for r in range(n_enc):
            fila_docx = docx_tabla.rows[r]
            tr_pr = fila_docx._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:tblHeader"))
            tr_pr.append(OxmlElement("w:cantSplit"))
            for celda in fila_docx.cells:
                for parrafo in celda.paragraphs:
                    parrafo.paragraph_format.keep_with_next = True

        # Combinar celdas (SPAN)
        for cmd in tabla._spanCmds:
            (sc, sr), (ec, er) = cmd[1], cmd[2]
            sc, ec = self._norm(sc, n_cols), self._norm(ec, n_cols)
            sr, er = self._norm(sr, n_filas), self._norm(er, n_filas)
            if (sr, sc) != (er, ec):
                docx_tabla.cell(sr, sc).merge(docx_tabla.cell(er, ec))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def _llenar_celda(self, celda, contenido, tam, texto_claro):
        parrafo = celda.paragraphs[0]
        parrafo.paragraph_format.space_after = Pt(0)
        color = RGBColor(0xFF, 0xFF, 0xFF) if texto_claro else TEXTO
        if isinstance(contenido, Paragraph):
            nombre = contenido.style.name
            if contenido.style.alignment == 1:
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            negrita = "bold" in getattr(contenido.style, "fontName", "").lower() or "Bold" in contenido.style.fontName
            _agregar_runs(parrafo, contenido.text, tam, color if texto_claro else NAVY if negrita else TEXTO,
                          negrita_base=negrita)
        elif isinstance(contenido, Table):
            textos = []
            for fila in contenido._cellvalues:
                for x in fila:
                    if isinstance(x, Paragraph):
                        textos.append(re.sub(r"<[^>]+>", "", x.text))
            # cuadro de color (punto) sin texto: se pasa su color a la celda
            for cmd in contenido._bkgrndcmds:
                if cmd[0] == "BACKGROUND" and hasattr(cmd[3], "red") and not textos:
                    _sombrear(celda, _hex(cmd[3]))
            _agregar_runs(parrafo, "  ".join(textos), tam, color)
        elif isinstance(contenido, str) and contenido:
            _agregar_runs(parrafo, contenido, tam, color)

    def guardar(self, ruta):
        self.doc.save(str(ruta))


def exportar(anio, salida=None):
    datos = g._datos_del_anio(anio)
    g._registrar_fuentes()
    estilos = g._estilos()
    story = g._construir_story(anio, estilos, datos, con_indice=False)

    exp = Exportador(anio)
    exp.portada(datos)
    for flowable in story[1:]:  # el primer elemento es el salto de página tras la portada
        exp.agregar(flowable)

    salida = Path(salida) if salida else g.BASE_DIR / f"Informe Anual de Calidad del Aire - Jalisco {anio}.docx"
    exp.guardar(salida)
    return salida


def _main():
    parser = argparse.ArgumentParser(description="Exporta el informe a Word (.docx).")
    parser.add_argument("--anio", type=int, required=True)
    parser.add_argument("--salida", type=str, default=None)
    args = parser.parse_args()
    print("Word generado en:", exportar(args.anio, args.salida))


if __name__ == "__main__":
    _main()
